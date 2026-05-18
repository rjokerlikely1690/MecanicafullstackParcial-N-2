from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_VERTICAL_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


BASE = Path(__file__).resolve().parent
OUT_PPTX = BASE / "Presentacion_Defensa_AutoMax.pptx"
OUT_MD = BASE / "Guion_Defensa_AutoMax.md"
TESTS_EVIDENCE_IMG = BASE / "evidencia_tests.png"


SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
TOP_BAND_H = Inches(0.72)
FOOTER_Y = Inches(7.06)

NAVY = RGBColor(15, 39, 74)
BLUE = RGBColor(37, 99, 235)
SKY = RGBColor(224, 242, 254)
ORANGE = RGBColor(249, 115, 22)
GREEN = RGBColor(22, 163, 74)
LIME = RGBColor(236, 253, 245)
SAND = RGBColor(255, 247, 237)
ROSE = RGBColor(255, 241, 242)
GRAY_BG = RGBColor(248, 250, 252)
CARD_BG = RGBColor(255, 255, 255)
TEXT = RGBColor(30, 41, 59)
MUTED = RGBColor(100, 116, 139)
BORDER = RGBColor(203, 213, 225)
WHITE = RGBColor(255, 255, 255)


def add_full_bg(slide, color=WHITE):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_textbox(
    slide,
    left,
    top,
    width,
    height,
    lines,
    *,
    font_size=18,
    color=TEXT,
    bold=False,
    align=PP_ALIGN.LEFT,
    font_name="Arial",
    margin=0.05,
    italic=False,
    vertical_anchor=MSO_VERTICAL_ANCHOR.TOP,
    space_after=3,
):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = vertical_anchor
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.clear()
    if isinstance(lines, str):
        lines = lines.split("\n")
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.space_after = Pt(space_after)
        p.font.name = font_name
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.bold = bold if idx == 0 else False
        p.font.italic = italic
    return box


def add_header(slide, title, subtitle="AutoMax | Proyecto Fullstack", slide_no=None):
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, TOP_BAND_H)
    band.fill.solid()
    band.fill.fore_color.rgb = NAVY
    band.line.fill.background()

    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, TOP_BAND_H - Inches(0.07), SLIDE_W, Inches(0.07))
    accent.fill.solid()
    accent.fill.fore_color.rgb = ORANGE
    accent.line.fill.background()

    add_textbox(
        slide,
        0.55,
        0.12,
        10.0,
        0.38,
        [title],
        font_size=22,
        color=WHITE,
        bold=True,
        margin=0.0,
    )
    add_textbox(
        slide,
        10.8,
        0.16,
        2.1,
        0.3,
        [subtitle],
        font_size=10,
        color=RGBColor(219, 234, 254),
        bold=False,
        align=PP_ALIGN.RIGHT,
        margin=0.0,
    )
    add_textbox(
        slide,
        0.55,
        6.95,
        11.4,
        0.22,
        ["AutoMax · proyecto fullstack"],
        font_size=9,
        color=MUTED,
        margin=0.0,
    )
    if slide_no is not None:
        add_textbox(
            slide,
            12.15,
            6.92,
            0.75,
            0.22,
            [str(slide_no)],
            font_size=10,
            color=MUTED,
            bold=True,
            align=PP_ALIGN.RIGHT,
            margin=0.0,
        )


def add_card(
    slide,
    left,
    top,
    width,
    height,
    title,
    body_lines,
    accent=NAVY,
    fill=CARD_BG,
    title_color=WHITE,
    body_align=PP_ALIGN.LEFT,
    body_vertical_anchor=MSO_VERTICAL_ANCHOR.TOP,
    body_font_size=14,
    body_margin=0.0,
    body_space_after=3,
    body_top_offset=0.58,
    body_height_offset=0.7,
    title_align=PP_ALIGN.LEFT,
):
    panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = fill
    panel.line.color.rgb = BORDER
    panel.line.width = Pt(1)

    header_h = 0.43
    header = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(header_h)
    )
    header.fill.solid()
    header.fill.fore_color.rgb = accent
    header.line.fill.background()

    add_textbox(
        slide,
        left + 0.15,
        top + 0.06,
        width - 0.3,
        0.22,
        [title],
        font_size=15,
        color=title_color,
        bold=True,
        margin=0.0,
        align=title_align,
    )
    add_textbox(
        slide,
        left + 0.15,
        top + body_top_offset,
        width - 0.3,
        height - body_height_offset,
        body_lines,
        font_size=body_font_size,
        color=TEXT,
        margin=body_margin,
        align=body_align,
        vertical_anchor=body_vertical_anchor,
        space_after=body_space_after,
    )


def add_section_title(slide, title, subtitle):
    add_textbox(
        slide,
        0.75,
        1.0,
        11.9,
        0.55,
        [title],
        font_size=27,
        color=NAVY,
        bold=True,
        margin=0.0,
    )
    add_textbox(
        slide,
        0.77,
        1.56,
        11.8,
        0.35,
        [subtitle],
        font_size=13,
        color=MUTED,
        italic=False,
        margin=0.0,
    )


def add_panel_bg(slide, left, top, width, height, fill=GRAY_BG):
    panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = fill
    panel.line.color.rgb = BORDER
    panel.line.width = Pt(1)
    return panel


def add_test_callout(slide, left, top, width, height, tag, accent, summary):
    add_panel_bg(slide, left, top, width, height, fill=WHITE)
    add_small_label(slide, left + 0.1, top + 0.08, 1.2, 0.22, tag, accent)
    add_textbox(
        slide,
        left + 1.42,
        top + 0.05,
        width - 1.54,
        height - 0.1,
        [summary],
        font_size=7.8,
        color=TEXT,
        margin=0.0,
        vertical_anchor=MSO_VERTICAL_ANCHOR.MIDDLE,
    )


def add_text_lines(slide, left, top, width, height, lines, font_size=18, color=TEXT, bold=False, margin=0.0):
    return add_textbox(
        slide,
        left,
        top,
        width,
        height,
        lines,
        font_size=font_size,
        color=color,
        bold=bold,
        margin=margin,
    )


def add_small_label(slide, left, top, width, height, label, fill, text_color=WHITE):
    label_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    label_box.fill.solid()
    label_box.fill.fore_color.rgb = fill
    label_box.line.fill.background()
    add_textbox(
        slide,
        left + 0.02,
        top + 0.02,
        width - 0.04,
        height - 0.04,
        [label],
        font_size=10,
        color=text_color,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0.0,
    )


def add_arrow(slide, left, top, width, height, fill=ORANGE):
    arrow = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = fill
    arrow.line.fill.background()
    return arrow


def add_connector_line(slide, x1, y1, x2, y2):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = BORDER
    line.line.width = Pt(2)
    return line


def write_markdown(notes):
    with OUT_MD.open("w", encoding="utf-8") as f:
        f.write("# Resumen de la presentación AutoMax\n\n")
        f.write("Resumen del contenido principal de la presentación.\n\n")
        for idx, item in enumerate(notes, start=1):
            f.write(f"## Slide {idx}. {item['title']}\n\n")
            f.write("- Contenido resumido en la diapositiva.\n\n")


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    prs.core_properties.title = "Presentación AutoMax"
    prs.core_properties.subject = "Branching, BFF, microservicios, patrones y pruebas"
    prs.core_properties.author = "Cursor"
    prs.core_properties.company = "AutoMax"

    blank = prs.slide_layouts[6]
    notes = []

    # Slide 1: portada
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, NAVY)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.22), SLIDE_H)
    accent.fill.solid()
    accent.fill.fore_color.rgb = ORANGE
    accent.line.fill.background()

    panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.85), Inches(1.05), Inches(7.0), Inches(4.9)
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor(17, 24, 39)
    panel.line.color.rgb = RGBColor(51, 65, 85)
    panel.line.width = Pt(1)

    add_textbox(
        slide,
        1.2,
        1.45,
        6.0,
        1.1,
        ["AutoMax", "Presentación del proyecto"],
        font_size=30,
        color=WHITE,
        bold=True,
        margin=0.0,
    )
    add_textbox(
        slide,
        1.2,
        2.7,
        6.1,
        1.0,
        [
            "Branching, arquitectura, patrones, pruebas y documentación.",
        ],
        font_size=18,
        color=RGBColor(219, 234, 254),
        margin=0.0,
    )
    add_textbox(
        slide,
        1.2,
        3.8,
        6.1,
        0.7,
        ["DSY1106 · Desarrollo Fullstack III"],
        font_size=16,
        color=RGBColor(147, 197, 253),
        bold=True,
        margin=0.0,
    )
    add_textbox(
        slide,
        1.2,
        4.35,
        6.1,
        0.7,
        ["Demo, arquitectura, Git, pruebas y documentación"],
        font_size=15,
        color=RGBColor(191, 219, 254),
        margin=0.0,
    )

    add_card(
        slide,
        8.35,
        1.25,
        4.0,
        4.55,
        "Ejes del proyecto",
        [
            "• Frontend NPM y consumo por API.",
            "• BFF / Gateway como entrada única.",
            "• Microservicios y responsabilidades separadas.",
            "• Branching con Git y resolución de conflictos.",
            "• Patrones, pruebas y documentación.",
        ],
        accent=BLUE,
        fill=GRAY_BG,
        title_color=WHITE,
    )
    add_textbox(
        slide,
        8.55,
        6.05,
        3.6,
        0.45,
        ["Demo, arquitectura, Git, pruebas y documentación."],
        font_size=11,
        color=RGBColor(226, 232, 240),
        italic=True,
        margin=0.0,
    )
    notes.append(
        {
            "title": "Portada",
            "notes": [
                "AutoMax resume una solución fullstack para un taller mecánico.",
                "La propuesta integra arquitectura, branching, patrones y pruebas.",
                "La documentación sostiene la trazabilidad técnica del proyecto.",
            ],
        }
    )

    # Slide 2: qué pide la pauta
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, WHITE)
    add_header(slide, "¿Qué pide la pauta?", "Criterios de entrega", slide_no=2)
    add_section_title(slide, "Requisitos de la pauta", "Frontend, BFF, microservicios y documentación.")
    add_card(
        slide,
        0.75,
        2.05,
        3.85,
        2.0,
        "Encargo grupal",
        [
            "• Frontend NPM.",
            "• Un BFF y al menos dos microservicios.",
            "• Arquetipos Maven para backend.",
            "• Todo versionado en GitHub.",
        ],
        accent=BLUE,
        fill=SKY,
        title_color=WHITE,
    )
    add_card(
        slide,
        4.77,
        2.05,
        3.85,
        2.0,
        "Documentación técnica",
        [
            "• Patrones y arquetipos.",
            "• Branching en Git.",
            "• README por componente.",
            "• `repositorios.txt` con enlaces finales.",
        ],
        accent=GREEN,
        fill=LIME,
        title_color=WHITE,
    )
    add_card(
        slide,
        8.79,
        2.05,
        3.8,
        2.0,
        "Evaluación oral",
        [
            "• 15 minutos en total.",
            "• Exposición individual.",
            "• Preguntas sobre patrones, branching y pruebas.",
            "• 70% del puntaje.",
        ],
        accent=ORANGE,
        fill=SAND,
        title_color=WHITE,
    )
    add_card(
        slide,
        0.75,
        4.45,
        11.85,
        1.5,
        "Criterio principal",
        [
            "Arquitectura, Git, documentación y pruebas respaldan la calidad de la solución.",
        ],
        accent=NAVY,
        fill=GRAY_BG,
        title_color=WHITE,
    )
    notes.append(
        {
            "title": "Qué pide la pauta",
            "notes": [
                "La entrega integra frontend NPM, BFF, microservicios, arquetipos Maven y control en GitHub.",
                "La documentación reúne patrones, branching, README por componente y `repositorios.txt`.",
                "La evaluación considera exposición individual y evidencia técnica de la solución.",
            ],
        }
    )

    # Slide 3: arquitectura general
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, WHITE)
    add_header(slide, "Arquitectura general de AutoMax", "Frontend, BFF, microservicios y datos", slide_no=3)
    add_section_title(slide, "Flujo de arquitectura", "Frontend, BFF, microservicios y datos.")

    # Flow boxes
    add_card(
        slide,
        0.55,
        2.15,
        2.95,
        1.75,
        "Frontend NPM",
        [
            "• React y Axios.",
            "• Rutas protegidas.",
            "• Consume una API única.",
        ],
        accent=BLUE,
        fill=SKY,
        title_color=WHITE,
    )
    add_arrow(slide, 3.63, 2.65, 0.45, 0.45, fill=ORANGE)
    add_card(
        slide,
        4.15,
        2.15,
        2.95,
        1.75,
        "BFF / Gateway",
        [
            "• Entrada única.",
            "• Adapta respuestas.",
            "• Centraliza seguridad.",
        ],
        accent=ORANGE,
        fill=SAND,
        title_color=WHITE,
    )
    add_arrow(slide, 7.23, 2.65, 0.45, 0.45, fill=GREEN)
    add_card(
        slide,
        7.75,
        2.15,
        2.95,
        1.75,
        "Microservicios",
        [
            "• Usuarios.",
            "• Vehículos.",
            "• Turnos y servicios.",
        ],
        accent=GREEN,
        fill=LIME,
        title_color=WHITE,
    )
    add_arrow(slide, 10.83, 2.65, 0.45, 0.45, fill=BLUE)
    add_card(
        slide,
        11.35,
        2.15,
        1.4,
        1.75,
        "Persistencia",
        [
            "• H2 en memoria.",
            "• MongoDB.",
        ],
        accent=BLUE,
        fill=GRAY_BG,
        title_color=WHITE,
    )
    add_card(
        slide,
        0.75,
        4.45,
        5.9,
        1.55,
        "Arquitectura integrada",
        [
            "El frontend conversa con una capa de entrada que distribuye la lógica hacia servicios especializados.",
        ],
        accent=NAVY,
        fill=GRAY_BG,
        title_color=WHITE,
    )
    add_card(
        slide,
        6.85,
        4.45,
        5.7,
        1.55,
        "Ventajas del diseño",
        [
            "La separación de responsabilidades y el control de cambios ordenan la solución.",
        ],
        accent=ORANGE,
        fill=SAND,
        title_color=WHITE,
    )
    notes.append(
        {
            "title": "Arquitectura general",
            "notes": [
                "El frontend React consume una API única mediante Axios y rutas protegidas.",
                "El gateway centraliza la entrada, adapta respuestas y concentra seguridad.",
                "Los microservicios separan usuarios, vehículos, turnos y servicios; H2 y Mongo respaldan los datos.",
            ],
        }
    )

    # Slide 4: frontend
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, WHITE)
    add_header(slide, "Frontend NPM", "Elementos principales del cliente web", slide_no=4)
    add_section_title(slide, "Arquitectura del frontend", "Organización por componentes y control de acceso.")
    add_card(
        slide,
        0.75,
        2.0,
        6.0,
        3.9,
        "Componentes clave",
        [
            "• React organiza la interfaz por componentes.",
            "• `package.json` concentra dependencias y scripts.",
            "• Axios centraliza las llamadas HTTP.",
            "• `AuthContext` guarda sesión y usuario.",
            "• `ProtectedRoute` controla accesos según rol.",
            "• Bootstrap y React-Bootstrap mantienen consistencia visual.",
        ],
        accent=BLUE,
        fill=SKY,
        title_color=WHITE,
    )
    add_card(
        slide,
        7.0,
        2.0,
        5.55,
        1.8,
        "Arquitectura del frontend",
        [
            "El frontend muestra datos, pide acciones y protege rutas sin mezclar lógica de negocio.",
        ],
        accent=GREEN,
        fill=LIME,
        title_color=WHITE,
    )
    add_card(
        slide,
        7.0,
        4.05,
        5.55,
        1.85,
        "Rol del frontend",
        [
            "El frontend consume la API, controla la navegación y mejora la experiencia del usuario.",
        ],
        accent=ORANGE,
        fill=SAND,
        title_color=WHITE,
    )
    notes.append(
        {
            "title": "Frontend NPM",
            "notes": [
                "El frontend NPM organiza la interfaz por componentes y control de acceso.",
                "`package.json`, `Axios`, `AuthContext` y `ProtectedRoute` estructuran el cliente.",
                "La lógica de negocio permanece en el backend.",
            ],
        }
    )

    # Slide 5: BFF / Gateway
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, WHITE)
    add_header(slide, "BFF / Gateway", "Entrada única, adaptación y seguridad", slide_no=5)
    add_section_title(slide, "Función del BFF", "Entrada única, adaptación y seguridad.")

    add_card(
        slide,
        0.75,
        2.05,
        2.85,
        1.8,
        "Frontend NPM",
        [
            "• Pide una acción.",
            "• No conoce detalles internos.",
        ],
        accent=BLUE,
        fill=SKY,
        title_color=WHITE,
    )
    add_arrow(slide, 3.8, 2.63, 0.45, 0.45, fill=ORANGE)
    add_card(
        slide,
        4.3,
        2.05,
        3.15,
        1.8,
        "BFF / Gateway",
        [
            "• Agrega datos.",
            "• Traduce respuestas.",
            "• Filtra y valida.",
        ],
        accent=ORANGE,
        fill=SAND,
        title_color=WHITE,
    )
    add_arrow(slide, 7.63, 2.63, 0.45, 0.45, fill=GREEN)
    add_card(
        slide,
        8.15,
        2.05,
        4.4,
        1.8,
        "Servicios específicos",
        [
            "• Usuarios.",
            "• Vehículos.",
            "• Turnos.",
            "• Catálogo de servicios.",
        ],
        accent=GREEN,
        fill=LIME,
        title_color=WHITE,
    )
    add_card(
        slide,
        0.75,
        4.3,
        5.9,
        1.7,
        "Comunicación centralizada",
        [
            "El BFF consolida peticiones del frontend; por ejemplo, al registrar un turno, la solicitud llega a una sola API.",
        ],
        accent=NAVY,
        fill=GRAY_BG,
        title_color=WHITE,
    )
    add_card(
        slide,
        6.9,
        4.3,
        5.65,
        1.7,
        "Gateway como entrada",
        [
            "El gateway valida cada solicitud; por ejemplo, una consulta de vehículos va a `vehiculos-service`.",
        ],
        accent=BLUE,
        fill=SKY,
        title_color=WHITE,
    )
    notes.append(
        {
            "title": "BFF / Gateway",
            "notes": [
                "El BFF o gateway evita exponer los servicios internos al cliente.",
                "Agrega respuestas, traduce formatos y centraliza validaciones.",
                "El frontend consume una API única y más clara.",
            ],
        }
    )

    # Slide 6: microservices
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, WHITE)
    add_header(slide, "Microservicios", "Servicios separados por dominio", slide_no=6)
    add_section_title(slide, "Separación por dominios", "Usuarios, vehículos, turnos y catálogo.")

    add_card(
        slide,
        0.6,
        2.0,
        3.0,
        1.9,
        "usuarios-service",
        [
            "• Registro y login.",
            "• Control de usuarios y roles.",
            "• Base para autenticación.",
        ],
        accent=BLUE,
        fill=SKY,
        title_color=WHITE,
    )
    add_card(
        slide,
        3.75,
        2.0,
        3.0,
        1.9,
        "vehiculos-service",
        [
            "• Gestión de autos.",
            "• Estado del vehículo.",
            "• Datos asociados al cliente.",
        ],
        accent=GREEN,
        fill=LIME,
        title_color=WHITE,
    )
    add_card(
        slide,
        6.9,
        2.0,
        3.0,
        1.9,
        "turnos-service",
        [
            "• Reserva de citas.",
            "• Fechas y horarios.",
            "• Relación con vehículo y usuario.",
        ],
        accent=ORANGE,
        fill=SAND,
        title_color=WHITE,
    )
    add_card(
        slide,
        10.05,
        2.0,
        2.65,
        1.9,
        "Autenticación",
        [
            "• Login y registro.",
            "• Base para acceso seguro.",
            "• Soporte de acceso.",
        ],
        accent=NAVY,
        fill=GRAY_BG,
        title_color=WHITE,
    )
    add_card(
        slide,
        0.75,
        4.35,
        5.9,
        1.7,
        "Separación de dominios",
        [
            "Cada microservicio resuelve un problema puntual. Eso permite cambiar una parte sin romper todo el sistema.",
        ],
        accent=GREEN,
        fill=LIME,
        title_color=WHITE,
    )
    add_card(
        slide,
        6.9,
        4.35,
        5.65,
        1.7,
        "Arquitectura evolutiva",
        [
            "La separación de dominios define la comunicación entre servicios y favorece el trabajo independiente.",
        ],
        accent=BLUE,
        fill=SKY,
        title_color=WHITE,
    )
    notes.append(
        {
            "title": "Microservicios",
            "notes": [
                "Los servicios distribuyen el dominio en usuarios, vehículos, turnos y catálogo.",
                "Cada servicio evoluciona de forma independiente.",
                "La separación facilita mantenimiento y escalabilidad.",
            ],
        }
    )

    # Slide 7: design patterns
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, WHITE)
    add_header(slide, "Patrones de diseño", "Persistencia, negocio y acceso", slide_no=7)
    add_section_title(slide, "Aplicación de patrones", "Persistencia, negocio, acceso y comportamiento.")
    add_card(
        slide,
        0.7,
        2.0,
        3.0,
        1.9,
        "Repository",
        [
            "• Separa persistencia.",
            "• Encapsula consultas.",
            "• Facilita pruebas.",
        ],
        accent=BLUE,
        fill=SKY,
        title_color=WHITE,
    )
    add_card(
        slide,
        3.85,
        2.0,
        3.0,
        1.9,
        "Service Layer",
        [
            "• Reglas de negocio.",
            "• Orquesta operaciones.",
            "• Evita lógica en controladores.",
        ],
        accent=GREEN,
        fill=LIME,
        title_color=WHITE,
    )
    add_card(
        slide,
        7.0,
        2.0,
        3.0,
        1.9,
        "Strategy / Guards",
        [
            "• Cambia el comportamiento.",
            "• Ejemplo: rutas protegidas.",
            "• Reglas por rol o estado.",
        ],
        accent=ORANGE,
        fill=SAND,
        title_color=WHITE,
    )
    add_card(
        slide,
        10.15,
        2.0,
        2.5,
        1.9,
        "DI / Singleton",
        [
            "• Spring gestiona beans.",
            "• Reutilización segura.",
        ],
        accent=NAVY,
        fill=GRAY_BG,
        title_color=WHITE,
    )
    add_card(
        slide,
        0.75,
        4.35,
        12.0,
        1.65,
        "Relación entre patrones",
        [
            "Repository guarda el acceso a datos, Service Layer concentra la lógica, Strategy cambia comportamientos según el contexto y la inyección de dependencias evita acoplamiento fuerte.",
        ],
        accent=NAVY,
        fill=GRAY_BG,
        title_color=WHITE,
    )
    notes.append(
        {
            "title": "Patrones de diseño",
            "notes": [
                "`Repository` separa persistencia y consultas.",
                "`Service Layer` concentra reglas de negocio.",
                "`Strategy`, `Route Guards` e inyección de dependencias resuelven variaciones de comportamiento y acceso.",
            ],
        }
    )

    # Slide 8: arquetipos Maven
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, WHITE)
    add_header(slide, "Arquetipos Maven", "Estandarización del backend", slide_no=8)
    add_section_title(slide, "Estructura y repetibilidad", "Arquetipos Maven y proyectos consistentes.")
    add_card(
        slide,
        0.75,
        2.05,
        2.85,
        1.75,
        "Inicialización",
        [
            "• Usar un arquetipo base.",
            "• Crear estructura estándar.",
        ],
        accent=BLUE,
        fill=SKY,
        title_color=WHITE,
    )
    add_arrow(slide, 3.78, 2.58, 0.45, 0.45, fill=ORANGE)
    add_card(
        slide,
        4.3,
        2.05,
        2.85,
        1.75,
        "Configuración",
        [
            "• `pom.xml`.",
            "• Dependencias y plugins.",
        ],
        accent=ORANGE,
        fill=SAND,
        title_color=WHITE,
    )
    add_arrow(slide, 7.33, 2.58, 0.45, 0.45, fill=GREEN)
    add_card(
        slide,
        7.85,
        2.05,
        2.85,
        1.75,
        "Estructura base",
        [
            "• `src/main/java`.",
            "• `src/main/resources`.",
        ],
        accent=GREEN,
        fill=LIME,
        title_color=WHITE,
    )
    add_arrow(slide, 10.88, 2.58, 0.45, 0.45, fill=BLUE)
    add_card(
        slide,
        11.4,
        2.05,
        1.2,
        1.75,
        "Documentación",
        [
            "• README.",
        ],
        accent=NAVY,
        fill=GRAY_BG,
        title_color=WHITE,
    )
    add_card(
        slide,
        0.75,
        4.35,
        12.0,
        1.75,
        "Ventajas",
        [
            "Los arquetipos dan una estructura común, reducen errores y aceleran la creación de servicios.",
        ],
        accent=BLUE,
        fill=SKY,
        title_color=WHITE,
    )
    notes.append(
        {
            "title": "Arquetipos Maven",
            "notes": [
                "Los arquetipos estandarizan la creación de cada backend.",
                "`pom.xml`, carpetas estándar y README mantienen coherencia entre servicios.",
                "La estructura repetible reduce errores y acelera el arranque.",
            ],
        }
    )

    # Slide 9: branching
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, WHITE)
    add_header(slide, "Branching en Git", "Flujo de integración", slide_no=9)
    add_section_title(slide, "Estructura de ramas", "main, develop y feature/*.")

    add_connector_line(slide, 6.7, 3.2, 6.7, 3.35)
    add_connector_line(slide, 6.7, 4.55, 2.425, 4.75)
    add_connector_line(slide, 6.7, 4.55, 6.625, 4.75)
    add_connector_line(slide, 6.7, 4.55, 10.825, 4.75)

    add_card(
        slide,
        4.75,
        2.0,
        3.9,
        1.2,
        "main",
        ["Rama estable"],
        accent=NAVY,
        fill=GRAY_BG,
        title_color=WHITE,
        body_align=PP_ALIGN.CENTER,
        body_vertical_anchor=MSO_VERTICAL_ANCHOR.MIDDLE,
        body_font_size=11,
        body_space_after=0,
        body_top_offset=0.49,
        body_height_offset=0.5,
        title_align=PP_ALIGN.CENTER,
    )
    add_card(
        slide,
        4.75,
        3.35,
        3.9,
        1.2,
        "develop",
        ["Integra cambios del equipo"],
        accent=BLUE,
        fill=SKY,
        title_color=WHITE,
        body_align=PP_ALIGN.CENTER,
        body_vertical_anchor=MSO_VERTICAL_ANCHOR.MIDDLE,
        body_font_size=11,
        body_space_after=0,
        body_top_offset=0.49,
        body_height_offset=0.5,
        title_align=PP_ALIGN.CENTER,
    )
    add_card(
        slide,
        0.75,
        4.75,
        3.35,
        1.0,
        "feature/login",
        ["Cambios aislados"],
        accent=GREEN,
        fill=LIME,
        title_color=WHITE,
        body_align=PP_ALIGN.CENTER,
        body_vertical_anchor=MSO_VERTICAL_ANCHOR.MIDDLE,
        body_font_size=11,
        body_space_after=0,
        body_top_offset=0.48,
        body_height_offset=0.5,
        title_align=PP_ALIGN.CENTER,
    )
    add_card(
        slide,
        4.95,
        4.75,
        3.35,
        1.0,
        "feature/bff",
        ["Endpoint y pruebas"],
        accent=ORANGE,
        fill=SAND,
        title_color=WHITE,
        body_align=PP_ALIGN.CENTER,
        body_vertical_anchor=MSO_VERTICAL_ANCHOR.MIDDLE,
        body_font_size=11,
        body_space_after=0,
        body_top_offset=0.48,
        body_height_offset=0.5,
        title_align=PP_ALIGN.CENTER,
    )
    add_card(
        slide,
        9.15,
        4.75,
        3.35,
        1.0,
        "feature/services",
        ["Microservicios"],
        accent=BLUE,
        fill=GRAY_BG,
        title_color=WHITE,
        body_align=PP_ALIGN.CENTER,
        body_vertical_anchor=MSO_VERTICAL_ANCHOR.MIDDLE,
        body_font_size=11,
        body_space_after=0,
        body_top_offset=0.48,
        body_height_offset=0.5,
        title_align=PP_ALIGN.CENTER,
    )
    add_card(
        slide,
        0.75,
        6.0,
        5.8,
        1.1,
        "Integración colaborativa",
        [
            "branch → commit → PR → revisión → merge",
        ],
        accent=GREEN,
        fill=LIME,
        title_color=WHITE,
        body_align=PP_ALIGN.CENTER,
        body_vertical_anchor=MSO_VERTICAL_ANCHOR.MIDDLE,
        body_font_size=11,
        body_space_after=0,
        body_top_offset=0.48,
        body_height_offset=0.5,
        title_align=PP_ALIGN.CENTER,
    )
    add_card(
        slide,
        6.75,
        6.0,
        5.95,
        1.1,
        "Resolución de conflictos",
        [
            "Se resuelven en la rama y se validan antes del merge.",
        ],
        accent=ORANGE,
        fill=SAND,
        title_color=WHITE,
        body_align=PP_ALIGN.CENTER,
        body_vertical_anchor=MSO_VERTICAL_ANCHOR.MIDDLE,
        body_font_size=11,
        body_space_after=0,
        body_top_offset=0.48,
        body_height_offset=0.5,
        title_align=PP_ALIGN.CENTER,
    )
    notes.append(
        {
            "title": "Branching",
            "notes": [
                "`main` concentra versiones estables.",
                "`develop` integra cambios del equipo.",
                "`feature/*` permite trabajo paralelo y trazabilidad mediante PR y merges.",
                "Los conflictos se resuelven antes de integrar cambios.",
            ],
        }
    )

    # Slide 10: buenas prácticas y pruebas
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, WHITE)
    add_header(slide, "Buenas prácticas y pruebas", "Mantenimiento y validación de la solución", slide_no=10)
    add_section_title(slide, "Calidad del código", "Mantenimiento y validación de la solución.")
    add_card(
        slide,
        0.75,
        2.05,
        5.85,
        2.9,
        "Buenas prácticas",
        [
            "• Separación de responsabilidades.",
            "• Nombres claros en clases y métodos.",
            "• README por componente.",
            "• Variables de entorno y configuración externa.",
            "• Código limpio y estructura consistente.",
        ],
        accent=BLUE,
        fill=SKY,
        title_color=WHITE,
    )
    add_card(
        slide,
        6.85,
        2.05,
        5.8,
        2.9,
        "Pruebas",
        [
            "• Unit tests para servicios y lógica clave.",
            "• Mock de dependencias cuando corresponda.",
            "• Validación de endpoints importantes.",
            "• Cobertura de rutas críticas.",
            "• Evidencia de ejecución y resultados.",
        ],
        accent=GREEN,
        fill=LIME,
        title_color=WHITE,
    )
    add_card(
        slide,
        0.75,
        5.25,
        11.9,
        1.0,
        "Valor de las pruebas",
        [
            "Las pruebas no están solo para aprobar: sirven para evitar regresiones y demostrar que la arquitectura elegida realmente ayuda a mantener el sistema.",
        ],
        accent=ORANGE,
        fill=SAND,
        title_color=WHITE,
    )
    notes.append(
        {
            "title": "Buenas prácticas y pruebas",
            "notes": [
                "La separación de responsabilidades, los nombres claros y la configuración ordenada facilitan el mantenimiento.",
                "Los tests unitarios y la validación de endpoints cubren escenarios críticos.",
                "La cobertura reduce regresiones y respalda la calidad de los cambios.",
            ],
        }
    )

    # Slide 11: resumen técnico
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, WHITE)
    add_header(slide, "Síntesis técnica", "Arquitectura, Git y calidad", slide_no=11)
    add_section_title(slide, "Integración de componentes", "Frontend, backend y calidad.")
    add_card(
        slide,
        0.75,
        2.1,
        3.8,
        1.45,
        "Frontend NPM",
        [
            "• React y componentización.",
            "• Axios y autenticación.",
            "• ProtectedRoute y navegación segura.",
        ],
        accent=BLUE,
        fill=SKY,
        title_color=WHITE,
    )
    add_card(
        slide,
        4.8,
        2.1,
        3.8,
        1.45,
        "Backend y microservicios",
        [
            "• Gateway como entrada única.",
            "• Usuarios, vehículos, turnos y servicios.",
            "• Separación por dominios.",
        ],
        accent=ORANGE,
        fill=SAND,
        title_color=WHITE,
    )
    add_card(
        slide,
        8.85,
        2.1,
        3.8,
        1.45,
        "Git y calidad",
        [
            "• Branching en Git.",
            "• Pruebas unitarias.",
            "• Buenas prácticas y documentación.",
        ],
        accent=GREEN,
        fill=LIME,
        title_color=WHITE,
    )
    add_card(
        slide,
        0.75,
        4.0,
        12.0,
        1.6,
        "Integración del sistema",
        [
            "Frontend, backend y calidad trabajan juntos para mantener una solución ordenada y trazable.",
        ],
        accent=NAVY,
        fill=GRAY_BG,
        title_color=WHITE,
    )
    notes.append(
        {
            "title": "Resumen técnico",
            "notes": [
                "La solución integra frontend, backend y prácticas de calidad.",
                "Branching, pruebas y documentación sostienen la trazabilidad del proyecto.",
                "La arquitectura resultante es ordenada y verificable.",
            ],
        }
    )

    # Slide 12: evidencias y repositorios
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, WHITE)
    add_header(slide, "Documentación y trazabilidad", "Archivos y repositorios del proyecto", slide_no=12)
    add_section_title(slide, "Inventario documental", "Documentos técnicos y repositorios.")
    add_card(
        slide,
        0.75,
        2.05,
        5.9,
        3.25,
        "Documentos técnicos",
        [
            "• `EVALUACION_PARCIAL_2_ESTUDIANTE.pdf`: pauta base y criterios.",
            "• README por componente: ejecución, puertos y dependencias.",
            "• Patrones, arquetipos y branching: decisiones técnicas.",
            "• `repositorios.txt`: URLs finales verificables.",
            "• ERS, manual, API y testing: respaldo documental.",
        ],
        accent=BLUE,
        fill=SKY,
        title_color=WHITE,
        body_font_size=12.2,
        body_space_after=2,
    )
    add_card(
        slide,
        6.9,
        2.05,
        5.65,
        3.25,
        "Repositorios del proyecto",
        [
            "• `Mecanica-ev3`: repo principal consolidado.",
            "• `01_Monolito/`: app real ejecutable.",
            "• `02_Microservicios/`: arquitectura del informe.",
            "• `03_Documentacion/`: entregables y evidencia.",
            "• `Djangofullstacktrabajomecanica/`: referencia comparativa.",
        ],
        accent=GREEN,
        fill=LIME,
        title_color=WHITE,
        body_font_size=12.2,
        body_space_after=2,
    )
    add_card(
        slide,
        0.75,
        5.55,
        11.9,
        0.8,
        "Trazabilidad final",
        [
            "La trazabilidad final debe usar enlaces reales y verificables.",
        ],
        accent=ORANGE,
        fill=SAND,
        title_color=WHITE,
    )
    notes.append(
        {
            "title": "Evidencias y enlaces",
            "notes": [
                "La documentación técnica reúne la pauta, el informe, los requisitos, el manual y la cobertura.",
                "Los repositorios separan monolito, microservicios y documentación.",
                "La carpeta consolida el recorrido completo del proyecto.",
            ],
        }
    )

    # Slide 13: validaciones técnicas
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, WHITE)
    add_header(slide, "Validaciones técnicas", "BFF, branching y pruebas", slide_no=13)
    add_section_title(slide, "Conclusiones de la solución", "Arquitectura, integración y calidad.")
    add_card(
        slide,
        0.75,
        2.05,
        3.8,
        2.0,
        "BFF / Gateway",
        [
            "Centraliza la comunicación del frontend con el backend y reduce acoplamiento.",
        ],
        accent=BLUE,
        fill=SKY,
        title_color=WHITE,
        body_font_size=12.1,
        body_space_after=2,
    )
    add_card(
        slide,
        4.8,
        2.05,
        3.8,
        2.0,
        "Branching",
        [
            "main, develop y feature/* organizan integración y trabajo paralelo.",
        ],
        accent=ORANGE,
        fill=SAND,
        title_color=WHITE,
    )
    add_card(
        slide,
        8.85,
        2.05,
        3.8,
        2.0,
        "Pruebas",
        [
            "Validan el comportamiento crítico y evitan regresiones.",
        ],
        accent=GREEN,
        fill=LIME,
        title_color=WHITE,
    )
    add_card(
        slide,
        0.75,
        4.55,
        12.0,
        1.3,
        "Cierre técnico",
        [
            "Arquitectura, control de versiones y pruebas sostienen la calidad del proyecto.",
        ],
        accent=NAVY,
        fill=GRAY_BG,
        title_color=WHITE,
    )
    notes.append(
        {
            "title": "Validaciones técnicas",
            "notes": [
                "El BFF centraliza la comunicación y protege el acceso al backend.",
                "Branching organiza integración y trabajo paralelo.",
                "Las pruebas validan comportamiento crítico y evitan regresiones.",
                "La calidad sostiene la entrega de AutoMax.",
            ],
        }
    )

    # Slide 14: evidencia de pruebas
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, WHITE)
    add_header(slide, "Evidencia de pruebas", "Salida real de Maven", slide_no=14)
    add_section_title(slide, "Validación automatizada", "Autenticación y flujo de órdenes.")
    add_card(
        slide,
        0.75,
        2.05,
        7.6,
        4.55,
        "Qué valida",
        [
            "• `AuthServiceTest`: registro, login y JWT para acceso protegido.",
            "• `OrdenServiceTest`: creación de turno y orden asociada.",
            "• `mvn test`: ejecuta la suite completa y confirma estabilidad.",
            "• La terminal real demuestra trazabilidad de la ejecución.",
        ],
        accent=BLUE,
        fill=SKY,
        title_color=WHITE,
    )
    add_panel_bg(slide, 8.65, 2.05, 3.9, 4.55, fill=GRAY_BG)
    add_small_label(slide, 8.9, 2.18, 1.55, 0.3, "Salida real de Maven", NAVY)
    add_test_callout(
        slide,
        8.9, 2.46, 3.38, 0.34,
        "AuthServiceTest",
        BLUE,
        "Login y token JWT.",
    )
    add_test_callout(
        slide,
        8.9, 2.86, 3.38, 0.34,
        "OrdenServiceTest",
        ORANGE,
        "Turno + orden.",
    )
    add_test_callout(
        slide,
        8.9, 3.26, 3.38, 0.34,
        "BUILD SUCCESS",
        GREEN,
        "Suite sin fallos.",
    )
    slide.shapes.add_picture(
        str(TESTS_EVIDENCE_IMG),
        Inches(9.02),
        Inches(3.66),
        Inches(3.05),
    )
    notes.append(
        {
            "title": "Evidencia de pruebas",
            "notes": [
                "La captura muestra la ejecución de AuthServiceTest y OrdenServiceTest con BUILD SUCCESS.",
                "La evidencia respalda el funcionamiento real del backend antes de la defensa.",
            ],
        }
    )

    prs.save(OUT_PPTX)
    write_markdown(notes)
    print(f"Presentación creada en: {OUT_PPTX}")
    print(f"Guion creado en: {OUT_MD}")


if __name__ == "__main__":
    main()
